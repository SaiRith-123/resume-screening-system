"""File text extraction: PDF (pdfplumber/pypdf) + DOCX (python-docx) with OCR fallback.

Never raises for a malformed file - returns an ExtractionResult describing the
outcome so the caller can mark a candidate FAILED with a meaningful status (spec §4, §23).
"""
from __future__ import annotations

import os
from dataclasses import dataclass

from app.core.logging import get_logger
from app.services.resume_parser.ocr import ocr_pdf

logger = get_logger(__name__)

MIN_MEANINGFUL_CHARS = 120


@dataclass
class ExtractionResult:
    text: str
    method: str  # native | ocr | failed
    page_count: int | None = None
    error: str | None = None

    @property
    def ok(self) -> bool:
        return self.method != "failed" and bool(self.text.strip())


def _extract_pdf(path: str) -> ExtractionResult:
    page_count: int | None = None
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception as exc:
        logger.warning("pdfplumber failed for %s: %s", path, exc)
    if len(text.strip()) < MIN_MEANINGFUL_CHARS:
        try:
            from pypdf import PdfReader

            reader = PdfReader(path)
            page_count = page_count or len(reader.pages)
            alt = "\n".join((pg.extract_text() or "") for pg in reader.pages)
            if len(alt.strip()) > len(text.strip()):
                text = alt
        except Exception as exc:
            logger.warning("pypdf failed for %s: %s", path, exc)
    if len(text.strip()) >= MIN_MEANINGFUL_CHARS:
        return ExtractionResult(text=text, method="native", page_count=page_count)
    # scanned / image PDF -> OCR
    ocr_text = ocr_pdf(path)
    if len(ocr_text.strip()) >= MIN_MEANINGFUL_CHARS:
        return ExtractionResult(text=ocr_text, method="ocr", page_count=page_count)
    if text.strip() or ocr_text.strip():
        # some text but below the threshold -> still return it, flagged native
        return ExtractionResult(
            text=(text or ocr_text), method="native", page_count=page_count,
            error="low_text_density",
        )
    return ExtractionResult(text="", method="failed", page_count=page_count,
                            error="empty_or_unreadable_pdf")


def _extract_docx(path: str) -> ExtractionResult:
    try:
        import docx  # python-docx

        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    except Exception as exc:
        logger.warning("python-docx failed for %s: %s", path, exc)
        return ExtractionResult(text="", method="failed", error="malformed_docx")
    if not text.strip():
        return ExtractionResult(text="", method="failed", error="empty_docx")
    return ExtractionResult(text=text, method="native")


def detect_kind(filename: str, content_type: str | None = None) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        return "pdf"
    if ext in (".docx", ".doc"):
        return "docx"
    ct = (content_type or "").lower()
    if "pdf" in ct:
        return "pdf"
    if "word" in ct or "officedocument" in ct:
        return "docx"
    return "unknown"


def extract_text(path: str, content_type: str | None = None) -> str:
    """Backwards-compatible helper returning only the extracted text."""
    return extract(path, content_type).text


def extract(path: str, content_type: str | None = None) -> ExtractionResult:
    kind = detect_kind(path, content_type)
    if kind == "pdf":
        return _extract_pdf(path)
    if kind == "docx":
        return _extract_docx(path)
    return ExtractionResult(text="", method="failed", error="unsupported_file_type")
